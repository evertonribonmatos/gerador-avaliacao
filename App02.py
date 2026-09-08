import os
import re
import json
import time
import tempfile
from io import BytesIO
from typing import List, Dict, Any, Optional, Tuple

import pdfplumber
import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from PIL import Image
import pytesseract
from docx import Document
from docx.table import _Cell, Table
from docx.shared import Inches
from docx.oxml.ns import qn

# ==========================================
# CONFIGURAÇÕES
# ==========================================
TOTAL_QUESTOES = 10
MODELO_PRINCIPAL = "openai/gpt-oss-120b"
LARGURA_IMAGEM_POLEGADAS = 4.8

TEMPERATURA_GERACAO = 0.55
MAX_TENTATIVAS_MODELO = 3
MAX_TENTATIVAS_POR_QUESTAO = 4

TAMANHO_MINIMO_CONTEXTO = 180
SIMILARIDADE_MAXIMA_PERMITIDA = 0.72

# Caminho do Tesseract OCR no Windows
CAMINHO_TESSERACT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
if os.path.exists(CAMINHO_TESSERACT):
    pytesseract.pytesseract.tesseract_cmd = CAMINHO_TESSERACT

# ==========================================
# SEGREDOS / CLIENTE IA
# ==========================================
load_dotenv()

def obter_groq_api_key() -> str:
    try:
        if "GROQ_API_KEY" in st.secrets:
            return st.secrets["GROQ_API_KEY"]
    except Exception:
        pass

    chave_env = os.getenv("GROQ_API_KEY")
    if chave_env:
        return chave_env

    raise ValueError(
        "GROQ_API_KEY não encontrada. Configure em st.secrets ou no arquivo .env."
    )

GROQ_API_KEY = obter_groq_api_key()
client = Groq(api_key=GROQ_API_KEY)

# ==========================================
# UTILITÁRIOS
# ==========================================
def normalizar_texto(texto: Any) -> str:
    if texto is None:
        return ""
    texto = str(texto).replace("\r", "\n")
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()

def extrair_json_de_texto(texto: str) -> Dict[str, Any]:
    texto = texto.strip()
    try:
        return json.loads(texto)
    except json.JSONDecodeError:
        pass

    match = re.search(r"\{.*\}", texto, re.DOTALL)
    if match:
        return json.loads(match.group(0))

    raise ValueError("Não foi possível extrair JSON válido da resposta da IA.")

def chamar_ia_com_retry(**kwargs):
    ultima_excecao = None
    for tentativa in range(MAX_TENTATIVAS_MODELO):
        try:
            return client.chat.completions.create(**kwargs)
        except Exception as e:
            ultima_excecao = e
            espera = 2 ** tentativa
            time.sleep(espera)
    raise ultima_excecao

def uploaded_file_para_bytes(uploaded_file) -> Optional[bytes]:
    if uploaded_file is None:
        return None
    uploaded_file.seek(0)
    return uploaded_file.read()

def tentar_ocr_em_imagem(imagem_bytes: Optional[bytes]) -> str:
    if not imagem_bytes:
        return ""

    try:
        imagem = Image.open(BytesIO(imagem_bytes))
        texto = pytesseract.image_to_string(imagem, lang="por+eng")
        return normalizar_texto(texto)
    except Exception:
        return ""

def tokenizar_simples(texto: str) -> List[str]:
    return re.findall(r"\w+", normalizar_texto(texto).lower())

def similaridade_textual_simples(a: str, b: str) -> float:
    a_tokens = set(tokenizar_simples(a))
    b_tokens = set(tokenizar_simples(b))

    if not a_tokens or not b_tokens:
        return 0.0

    inter = len(a_tokens & b_tokens)
    uniao = len(a_tokens | b_tokens)
    return inter / uniao if uniao else 0.0

def contem_termos_proibidos_sem_suporte(contexto: str, possui_imagem: bool) -> Optional[str]:
    texto = normalizar_texto(contexto).lower()

    termos_proibidos_gerais = [
        "tabela", "relatório", "relatorio", "gráfico", "grafico",
        "laudo", "planilha", "prontuário", "prontuario", "anexo",
        "dados da tabela", "conforme tabela", "segundo o relatório",
        "de acordo com o relatório", "com base no relatório"
    ]

    termos_visuais = [
        "figura", "imagem", "diagrama", "ilustração", "ilustracao",
        "esquema abaixo", "figura abaixo", "imagem abaixo"
    ]

    for termo in termos_proibidos_gerais:
        if termo in texto:
            return termo

    if not possui_imagem:
        for termo in termos_visuais:
            if termo in texto:
                return termo

    return None

def contem_generico_demais(contexto: str) -> bool:
    texto = normalizar_texto(contexto).lower()

    padroes_genericos = [
        "explique o que é",
        "defina",
        "conceitue",
        "cite",
        "liste",
        "o que é manutenção",
        "o que é motor trifásico"
    ]

    return any(p in texto for p in padroes_genericos)

def validar_bloom(bloom: str) -> str:
    bloom = normalizar_texto(bloom)
    if not bloom:
        return "Analisar"
    return bloom

# ==========================================
# EXTRAÇÃO DE CONTEÚDO BASE
# ==========================================
def extrair_texto_pdf_upload(uploaded_file) -> str:
    textos = []
    uploaded_file.seek(0)

    with pdfplumber.open(uploaded_file) as pdf:
        for pagina in pdf.pages:
            texto = pagina.extract_text() or ""
            textos.append(texto)

    texto_final = normalizar_texto("\n".join(textos))
    if not texto_final:
        raise ValueError("Não foi possível extrair texto do PDF enviado.")
    return texto_final

def extrair_texto_txt_upload(uploaded_file) -> str:
    uploaded_file.seek(0)
    conteudo_bytes = uploaded_file.read()
    try:
        texto = conteudo_bytes.decode("utf-8")
    except UnicodeDecodeError:
        texto = conteudo_bytes.decode("latin-1")

    texto = normalizar_texto(texto)
    if not texto:
        raise ValueError("O arquivo TXT enviado está vazio.")
    return texto

def obter_conteudo_base_upload(modo_conteudo: str, arquivo_base, conteudo_manual: str) -> str:
    if modo_conteudo == "Texto manual":
        conteudo = normalizar_texto(conteudo_manual)
        if not conteudo:
            raise ValueError("O conteúdo-base manual não pode ficar vazio.")
        return conteudo

    if arquivo_base is None:
        raise ValueError("Envie um documento-base (.pdf ou .txt).")

    nome = arquivo_base.name.lower()
    if nome.endswith(".pdf"):
        return extrair_texto_pdf_upload(arquivo_base)
    elif nome.endswith(".txt"):
        return extrair_texto_txt_upload(arquivo_base)
    else:
        raise ValueError("O documento-base deve ser .pdf ou .txt.")

# ==========================================
# PROMPTS
# ==========================================
def montar_resumo_questoes_anteriores(questoes_anteriores: List[Dict[str, Any]]) -> str:
    if not questoes_anteriores:
        return "Nenhuma questão anterior."

    blocos = []
    for q in questoes_anteriores:
        contexto = normalizar_texto(q.get("contexto", ""))
        contexto_curto = contexto[:500]
        blocos.append(
            f"Questão {q['numero']:02d} | tipo={q['tipo']} | bloom={q.get('bloom', '')} | resumo={contexto_curto}"
        )
    return "\n".join(blocos)

def montar_prompt_questao_unica(
    conteudo_base: str,
    dados_usuario: Dict[str, Any],
    questao_atual: Dict[str, Any],
    questoes_anteriores: List[Dict[str, Any]],
    feedback_erro: str = ""
) -> str:
    resumo_anteriores = montar_resumo_questoes_anteriores(questoes_anteriores)

    possui_imagem = "sim" if questao_atual.get("imagem_bytes") else "não"
    ocr = normalizar_texto(questao_atual.get("ocr_imagem", ""))

    bloco_imagem = ""
    if questao_atual.get("imagem_bytes"):
        bloco_imagem = f"""
IMAGEM ASSOCIADA À QUESTÃO:
- Existe imagem vinculada a esta questão.
- Texto extraído por OCR:
{ocr if ocr else "Nenhum texto legível foi extraído da imagem."}

REGRAS ESPECÍFICAS DA IMAGEM:
- O enunciado deve deixar claro que a resolução depende da observação da imagem apresentada.
- A questão deve explorar identificação, interpretação, diagnóstico, análise técnica ou tomada de decisão com base na imagem.
- Não diga que não consegue ver a imagem.
"""
    else:
        bloco_imagem = """
REGRAS ESPECÍFICAS:
- Esta questão NÃO possui imagem.
- Portanto, é proibido mencionar figura, imagem, diagrama, esquema visual, gráfico ou qualquer recurso visual inexistente.
"""

    bloco_tipo = ""
    if questao_atual["tipo"] == "objetiva":
        bloco_tipo = """
A questão deve ser OBJETIVA com:
- contexto completo
- comando claro
- 5 alternativas obrigatórias: A, B, C, D e E
- somente 1 alternativa correta
- alternativas técnicas plausíveis
- gabarito obrigatório
"""
    else:
        bloco_tipo = """
A questão deve ser DISCURSIVA com:
- contexto completo
- comando discursivo robusto
- exigência de resposta estruturada
- sem alternativas
- sem gabarito em letra
- exigir análise técnica, justificativa, procedimento, critérios, sequência lógica, diagnóstico, segurança ou proposta de solução
"""

    return f"""
Atue como especialista em elaboração de avaliações técnicas para educação profissional industrial.

OBJETIVO:
Gerar SOMENTE a Questão {questao_atual['numero']:02d}, com alto rigor técnico, redação profissional e nível difícil.

DADOS DA AVALIAÇÃO:
- Curso: {dados_usuario['curso']}
- Unidade Curricular: {dados_usuario['unidade_curricular']}
- Valor da avaliação: {dados_usuario['valor_avaliacao']}

CONFIGURAÇÃO DA QUESTÃO ATUAL:
- Número: {questao_atual['numero']}
- Tipo: {questao_atual['tipo']}
- Peso: {questao_atual['peso']}
- Imagem associada: {possui_imagem}

{bloco_imagem}

QUESTÕES JÁ GERADAS:
{resumo_anteriores}

INSTRUÇÕES CRÍTICAS:
1. Baseie-se EXCLUSIVAMENTE no conteúdo-base fornecido.
2. NÃO repita cenários, estruturas, comandos, redações, sintomas, contextos ou focos técnicos das questões já geradas.
3. Esta nova questão deve ser substancialmente diferente das anteriores.
4. Crie uma situação profissional plausível da área industrial.
5. O enunciado deve ser tecnicamente denso, contextualizado e completo.
6. Evite superficialidade.
7. Não faça pergunta meramente conceitual ou definicional.
8. Não invente tabela, relatório, gráfico, laudo, anexo, prontuário, planilha ou dados não fornecidos.
9. Se não houver imagem, não mencione figura, imagem, diagrama, esquema visual ou elemento gráfico.
10. Se houver imagem, o enunciado deve depender dela.
11. Não usar markdown.
12. Não usar crases.
13. Escrever em português brasileiro formal e técnico.
14. Priorizar verbos cognitivos compatíveis com aplicar, analisar, avaliar ou criar.
15. O enunciado deve, sempre que possível, envolver diagnóstico, manutenção, inspeção, proteção, ensaio, falha, operação, segurança, sequência lógica ou tomada de decisão.
16. Não gere enunciado genérico.
17. Não copie nem parafraseie a mesma questão anterior.
18. O texto deve ser autossuficiente, sem remeter a material inexistente.
19. Antes de responder, faça checagem interna para impedir repetição e referências indevidas.

{bloco_tipo}

FEEDBACK DE REJEIÇÕES ANTERIORES:
{feedback_erro if feedback_erro else "Nenhum."}

FORMATO DE SAÍDA:
Retorne SOMENTE JSON válido no seguinte formato:

{{
  "questao": {{
    "numero": {questao_atual['numero']},
    "tipo": "{questao_atual['tipo']}",
    "peso": "{questao_atual['peso']}",
    "contexto": "texto completo da questão",
    "alternativas": {{
      "A": "texto",
      "B": "texto",
      "C": "texto",
      "D": "texto",
      "E": "texto"
    }},
    "gabarito": "A",
    "bloom": "Analisar"
  }}
}}

Se a questão for discursiva:
- "alternativas" deve ser {{}}
- "gabarito" deve ser ""

CONTEÚDO-BASE:
{conteudo_base}
"""

# ==========================================
# VALIDAÇÃO DAS QUESTÕES
# ==========================================
def validar_estrutura_questao_unica(
    q: Dict[str, Any],
    questao_config: Dict[str, Any]
) -> Dict[str, Any]:
    numero_esperado = questao_config["numero"]
    tipo_esperado = questao_config["tipo"]

    if q.get("numero") != numero_esperado:
        raise ValueError(
            f"Questão retornada com número incorreto. Esperado: {numero_esperado}, recebido: {q.get('numero')}"
        )

    if q.get("tipo") != tipo_esperado:
        raise ValueError(
            f"Questão {numero_esperado}: tipo incorreto. Esperado: {tipo_esperado}, recebido: {q.get('tipo')}"
        )

    contexto = normalizar_texto(q.get("contexto", ""))
    if not contexto:
        raise ValueError(f"Questão {numero_esperado}: contexto vazio.")

    q["contexto"] = contexto
    q["peso"] = questao_config["peso"]
    q["bloom"] = validar_bloom(q.get("bloom", ""))

    if tipo_esperado == "objetiva":
        alternativas = q.get("alternativas")
        if not isinstance(alternativas, dict):
            raise ValueError(f"Questão {numero_esperado}: alternativas inválidas.")

        alternativas_normalizadas = {}
        for letra in ["A", "B", "C", "D", "E"]:
            texto_alt = normalizar_texto(alternativas.get(letra, ""))
            if not texto_alt:
                raise ValueError(f"Questão {numero_esperado}: alternativa {letra} ausente.")
            alternativas_normalizadas[letra] = texto_alt

        gabarito = normalizar_texto(q.get("gabarito", ""))
        if gabarito not in ["A", "B", "C", "D", "E"]:
            raise ValueError(f"Questão {numero_esperado}: gabarito inválido.")

        q["alternativas"] = alternativas_normalizadas
        q["gabarito"] = gabarito
    else:
        q["alternativas"] = {}
        q["gabarito"] = ""

    q["imagem_bytes"] = questao_config.get("imagem_bytes")
    q["imagem_nome"] = questao_config.get("imagem_nome", "")
    q["ocr_imagem"] = questao_config.get("ocr_imagem", "")

    return q

def validar_qualidade_questao_unica(
    questao: Dict[str, Any],
    questoes_anteriores: List[Dict[str, Any]]
) -> None:
    numero = questao["numero"]
    contexto = normalizar_texto(questao.get("contexto", ""))
    possui_imagem = bool(questao.get("imagem_bytes"))

    if len(contexto) < TAMANHO_MINIMO_CONTEXTO:
        raise ValueError(
            f"Questão {numero}: enunciado muito curto ou superficial."
        )

    termo_proibido = contem_termos_proibidos_sem_suporte(contexto, possui_imagem)
    if termo_proibido:
        raise ValueError(
            f"Questão {numero}: menciona recurso não fornecido ou inadequado: '{termo_proibido}'."
        )

    if contem_generico_demais(contexto):
        raise ValueError(
            f"Questão {numero}: enunciado excessivamente genérico."
        )

    if questao["tipo"] == "objetiva":
        for letra, alt in questao["alternativas"].items():
            if len(normalizar_texto(alt)) < 8:
                raise ValueError(
                    f"Questão {numero}: alternativa {letra} muito curta."
                )

    for anterior in questoes_anteriores:
        contexto_ant = normalizar_texto(anterior.get("contexto", ""))

        if contexto.lower() == contexto_ant.lower():
            raise ValueError(
                f"Questão {numero}: repetição literal da questão {anterior['numero']}."
            )

        sim = similaridade_textual_simples(contexto, contexto_ant)
        if sim > SIMILARIDADE_MAXIMA_PERMITIDA:
            raise ValueError(
                f"Questão {numero}: muito semelhante à questão {anterior['numero']} (similaridade={sim:.2f})."
            )

def validar_conjunto_final_questoes(questoes: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if len(questoes) != TOTAL_QUESTOES:
        raise ValueError(f"Devem existir exatamente {TOTAL_QUESTOES} questões ao final.")

    numeros = sorted([q["numero"] for q in questoes])
    if numeros != list(range(1, TOTAL_QUESTOES + 1)):
        raise ValueError("Numeração final das questões inválida.")

    questoes.sort(key=lambda x: x["numero"])

    for i, q in enumerate(questoes):
        validar_qualidade_questao_unica(q, questoes[:i])

    return questoes

# ==========================================
# GERAÇÃO DAS QUESTÕES COM IA
# ==========================================
def gerar_questao_unica_com_ia(
    conteudo_base: str,
    dados_usuario: Dict[str, Any],
    questao_config: Dict[str, Any],
    questoes_anteriores: List[Dict[str, Any]]
) -> Tuple[Dict[str, Any], str]:
    ultimo_erro = ""

    for tentativa in range(1, MAX_TENTATIVAS_POR_QUESTAO + 1):
        prompt = montar_prompt_questao_unica(
            conteudo_base=conteudo_base,
            dados_usuario=dados_usuario,
            questao_atual=questao_config,
            questoes_anteriores=questoes_anteriores,
            feedback_erro=ultimo_erro
        )

        completion = chamar_ia_com_retry(
            model=MODELO_PRINCIPAL,
            messages=[{"role": "user", "content": prompt}],
            temperature=TEMPERATURA_GERACAO,
            max_tokens=3500,
            response_format={"type": "json_object"}
        )

        resposta = completion.choices[0].message.content.strip()

        try:
            dados = extrair_json_de_texto(resposta)
            if "questao" not in dados or not isinstance(dados["questao"], dict):
                raise ValueError("JSON não contém a chave 'questao' corretamente.")

            questao = validar_estrutura_questao_unica(dados["questao"], questao_config)
            validar_qualidade_questao_unica(questao, questoes_anteriores)
            return questao, resposta

        except Exception as e:
            ultimo_erro = str(e)

    raise ValueError(
        f"Falha ao gerar a questão {questao_config['numero']:02d} após múltiplas tentativas. Último erro: {ultimo_erro}"
    )

def gerar_questoes_ia(conteudo_base: str, dados_usuario: Dict[str, Any]) -> Tuple[List[Dict[str, Any]], str, str]:
    questoes_geradas = []
    respostas_brutas = []

    for questao_config in dados_usuario["questoes"]:
        questao, resposta = gerar_questao_unica_com_ia(
            conteudo_base=conteudo_base,
            dados_usuario=dados_usuario,
            questao_config=questao_config,
            questoes_anteriores=questoes_geradas
        )
        questoes_geradas.append(questao)
        respostas_brutas.append({
            "numero": questao["numero"],
            "resposta_bruta": resposta
        })

    questoes_geradas = validar_conjunto_final_questoes(questoes_geradas)

    resposta_bruta_unificada = json.dumps(
        {"respostas_brutas": respostas_brutas},
        ensure_ascii=False,
        indent=2
    )

    return questoes_geradas, resposta_bruta_unificada, MODELO_PRINCIPAL

# ==========================================
# WORD - APOIO
# ==========================================
def texto_celula(cell: _Cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()

def limpar_celula(cell: _Cell) -> None:
    cell.text = ""
    if not cell.paragraphs:
        cell.add_paragraph()

def escrever_linhas_e_imagem_na_celula(cell: _Cell, linhas: List[str], imagem_bytes: Optional[bytes] = None) -> None:
    limpar_celula(cell)

    if not linhas and not imagem_bytes:
        return

    primeira_linha_escrita = False

    for linha in linhas:
        if not primeira_linha_escrita:
            cell.paragraphs[0].text = linha
            primeira_linha_escrita = True
        else:
            p = cell.add_paragraph()
            p.text = linha

    if imagem_bytes:
        if primeira_linha_escrita:
            cell.add_paragraph("")
        else:
            cell.paragraphs[0].text = ""

        p_img = cell.add_paragraph()
        run = p_img.add_run()
        run.add_picture(BytesIO(imagem_bytes), width=Inches(LARGURA_IMAGEM_POLEGADAS))

def preencher_campos_simples_em_tabelas(document: Document, dados_usuario: Dict[str, Any]) -> None:
    mapa_labels = {
        "curso:": dados_usuario["curso"],
        "unidade curricular:": dados_usuario["unidade_curricular"],
        "turma:": dados_usuario["turma"],
        "aluno:": dados_usuario["aluno"],
        "matrícula:": dados_usuario["matricula"],
        "matricula:": dados_usuario["matricula"],
        "valor da avaliação:": dados_usuario["valor_avaliacao"],
        "valor da avaliacao:": dados_usuario["valor_avaliacao"],
        "data:": dados_usuario["data"],
    }

    for table in document.tables:
        for row in table.rows:
            for idx, cell in enumerate(row.cells):
                txt = normalizar_texto(texto_celula(cell)).lower()
                if txt in mapa_labels:
                    valor = mapa_labels[txt]
                    if idx + 1 < len(row.cells):
                        row.cells[idx + 1].text = valor

def encontrar_tabela_questoes(document: Document) -> Optional[Table]:
    for table in document.tables:
        texto_total = " ".join(
            normalizar_texto(texto_celula(cell)).lower()
            for row in table.rows for cell in row.cells
        )

        if (
            ("questão" in texto_total or "questao" in texto_total)
            and "peso" in texto_total
            and "ponto obtido" in texto_total
        ):
            return table
    return None

def formatar_texto_questao(questao: Dict[str, Any]) -> List[str]:
    linhas = []

    bloom = normalizar_texto(questao.get("bloom", ""))
    contexto = normalizar_texto(questao.get("contexto", ""))

    if bloom:
        linhas.append(f"Nível cognitivo predominante: {bloom}.")
        linhas.append("")

    if contexto:
        linhas.extend(contexto.split("\n"))

    if questao["tipo"] == "objetiva":
        linhas.append("")
        for letra in ["A", "B", "C", "D", "E"]:
            alt = normalizar_texto(questao["alternativas"].get(letra, ""))
            linhas.append(f"({letra}) {alt}")

    return linhas

def permitir_altura_automatica_linha(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()

    for child in trPr.findall(qn("w:trHeight")):
        trPr.remove(child)

def preencher_tabela_questoes(document: Document, questoes: List[Dict[str, Any]]) -> None:
    tabela = encontrar_tabela_questoes(document)
    if not tabela:
        raise ValueError("Não foi possível localizar a tabela de questões no modelo Word.")

    mapa_questoes = {q["numero"]: q for q in questoes}

    i = 0
    while i < len(tabela.rows):
        row = tabela.rows[i]
        textos = [normalizar_texto(texto_celula(c)) for c in row.cells]
        textos_lower = [t.lower() for t in textos]

        if any(t in ["questão", "questao"] for t in textos_lower):
            numero_questao = None

            for t in textos:
                t_limpo = t.strip()
                if t_limpo.isdigit():
                    numero_questao = int(t_limpo)
                    break

            if numero_questao in mapa_questoes:
                q = mapa_questoes[numero_questao]

                for idx_c, txt in enumerate(textos_lower):
                    if txt == "peso" and idx_c + 1 < len(row.cells):
                        row.cells[idx_c + 1].text = str(q["peso"])

                if i + 1 < len(tabela.rows):
                    row_contexto = tabela.rows[i + 1]
                    permitir_altura_automatica_linha(row_contexto)

                    textos_contexto = [normalizar_texto(texto_celula(c)).lower() for c in row_contexto.cells]

                    idx_contexto = None
                    for idx_c, txt in enumerate(textos_contexto):
                        if txt == "contexto":
                            idx_contexto = idx_c
                            break

                    if idx_contexto is not None:
                        for c in row_contexto.cells[idx_contexto:]:
                            limpar_celula(c)
                        cell_destino = row_contexto.cells[idx_contexto].merge(row_contexto.cells[-1])
                    else:
                        for c in row_contexto.cells:
                            limpar_celula(c)
                        if len(row_contexto.cells) > 1:
                            cell_destino = row_contexto.cells[0].merge(row_contexto.cells[-1])
                        else:
                            cell_destino = row_contexto.cells[0]

                    linhas_questao = formatar_texto_questao(q)
                    escrever_linhas_e_imagem_na_celula(
                        cell_destino,
                        linhas_questao,
                        imagem_bytes=q.get("imagem_bytes")
                    )

        i += 1

# ==========================================
# GERAÇÃO DO DOCX FINAL
# ==========================================
def preencher_documento_word_em_memoria(
    modelo_docx_bytes: bytes,
    dados_usuario: Dict[str, Any],
    questoes: List[Dict[str, Any]]
) -> bytes:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp_in:
        tmp_in.write(modelo_docx_bytes)
        caminho_entrada = tmp_in.name

    try:
        document = Document(caminho_entrada)
        preencher_campos_simples_em_tabelas(document, dados_usuario)
        preencher_tabela_questoes(document, questoes)

        output = BytesIO()
        document.save(output)
        output.seek(0)
        return output.getvalue()
    finally:
        if os.path.exists(caminho_entrada):
            os.remove(caminho_entrada)

# ==========================================
# STREAMLIT UI
# ==========================================
st.set_page_config(page_title="Gerador de Avaliação Técnica", layout="wide")
st.title("Gerador de Avaliação Técnica")
st.write("Preencha os dados abaixo para gerar a avaliação e baixar o arquivo Word preenchido.")

with st.form("form_avaliacao"):
    st.subheader("Dados gerais")

    col1, col2 = st.columns(2)
    with col1:
        curso = st.text_input("Curso")
        unidade_curricular = st.text_input("Unidade Curricular")
        turma = st.text_input("Turma")
        aluno = st.text_input("Aluno")
    with col2:
        matricula = st.text_input("Matrícula")
        data = st.text_input("Data")
        valor_avaliacao = st.text_input("Valor da avaliação")

    modelo_docx = st.file_uploader("Modelo Word (.docx)", type=["docx"])

    st.subheader("Conteúdo-base")
    modo_conteudo = st.radio(
        "Forma de informar o conteúdo-base",
        ["Upload de arquivo (.pdf/.txt)", "Texto manual"]
    )

    arquivo_base = None
    conteudo_base_manual = ""

    if modo_conteudo == "Upload de arquivo (.pdf/.txt)":
        arquivo_base = st.file_uploader("Documento-base (.pdf ou .txt)", type=["pdf", "txt"])
    else:
        conteudo_base_manual = st.text_area(
            "Digite o conteúdo-base",
            height=250,
            placeholder="Cole aqui o conteúdo-base..."
        )

    st.subheader("Configuração das 10 questões")
    questoes_config = []

    for i in range(1, TOTAL_QUESTOES + 1):
        st.markdown(f"**Questão {i:02d}**")
        c1, c2 = st.columns(2)

        with c1:
            tipo = st.selectbox(
                f"Tipo da Questão {i:02d}",
                options=["objetiva", "discursiva"],
                key=f"tipo_{i}"
            )

        with c2:
            peso = st.text_input(
                f"Peso da Questão {i:02d}",
                value="1,0",
                key=f"peso_{i}"
            )

        imagem_questao = st.file_uploader(
            f"Imagem da Questão {i:02d} (opcional)",
            type=["png", "jpg", "jpeg"],
            key=f"imagem_{i}"
        )

        imagem_bytes = uploaded_file_para_bytes(imagem_questao)
        ocr_imagem = tentar_ocr_em_imagem(imagem_bytes)

        questoes_config.append({
            "numero": i,
            "tipo": tipo,
            "peso": peso,
            "imagem_bytes": imagem_bytes,
            "imagem_nome": imagem_questao.name if imagem_questao else "",
            "ocr_imagem": ocr_imagem
        })

    submitted = st.form_submit_button("Gerar avaliação")

if submitted:
    try:
        if modelo_docx is None:
            st.error("Envie o modelo Word (.docx).")
            st.stop()

        dados_usuario = {
            "curso": curso,
            "unidade_curricular": unidade_curricular,
            "turma": turma,
            "aluno": aluno,
            "matricula": matricula,
            "data": data,
            "valor_avaliacao": valor_avaliacao,
            "questoes": questoes_config
        }

        with st.spinner("Extraindo conteúdo-base..."):
            conteudo_base = obter_conteudo_base_upload(
                modo_conteudo=modo_conteudo,
                arquivo_base=arquivo_base,
                conteudo_manual=conteudo_base_manual
            )

        progresso = st.progress(0)
        status = st.empty()

        questoes_geradas = []
        respostas_brutas = []

        for idx, questao_config in enumerate(dados_usuario["questoes"], start=1):
            status.info(f"Gerando questão {idx:02d} de {TOTAL_QUESTOES}...")
            questao, resposta = gerar_questao_unica_com_ia(
                conteudo_base=conteudo_base,
                dados_usuario=dados_usuario,
                questao_config=questao_config,
                questoes_anteriores=questoes_geradas
            )
            questoes_geradas.append(questao)
            respostas_brutas.append({
                "numero": idx,
                "resposta_bruta": resposta
            })
            progresso.progress(idx / TOTAL_QUESTOES)

        status.info("Validando conjunto final das questões...")
        questoes = validar_conjunto_final_questoes(questoes_geradas)

        resposta_bruta_unificada = json.dumps(
            {"respostas_brutas": respostas_brutas},
            ensure_ascii=False,
            indent=2
        )

        with st.spinner("Preenchendo documento Word..."):
            docx_final_bytes = preencher_documento_word_em_memoria(
                modelo_docx_bytes=modelo_docx.getvalue(),
                dados_usuario=dados_usuario,
                questoes=questoes
            )

        status.empty()
        progresso.empty()

        st.success("Avaliação gerada com sucesso!")
        st.info(f"Modelo usado: {MODELO_PRINCIPAL}")

        with st.expander("Resposta bruta da IA"):
            st.text(resposta_bruta_unificada)

        with st.expander("Visualizar JSON gerado pela IA"):
            st.json({"questoes": questoes})

        with st.expander("Debug das questões geradas"):
            for q in questoes:
                st.write(
                    f"Questão {q['numero']} | tipo: {q['tipo']} | peso: {q['peso']} | "
                    f"imagem: {'sim' if q.get('imagem_bytes') else 'não'} | bloom: {q.get('bloom', '')}"
                )

                if q.get("ocr_imagem"):
                    st.write(f"OCR da imagem: {q['ocr_imagem'][:300]}")

                st.write("Contexto:")
                st.write(q["contexto"][:1200] + "..." if len(q["contexto"]) > 1200 else q["contexto"])

                if q["tipo"] == "objetiva":
                    st.write("Alternativas:")
                    for letra in ["A", "B", "C", "D", "E"]:
                        st.write(f"{letra}: {q['alternativas'].get(letra, '')}")
                    st.write(f"Gabarito: {q.get('gabarito', '')}")

                st.write("---")

        st.download_button(
            label="Baixar avaliação preenchida (.docx)",
            data=docx_final_bytes,
            file_name="avaliacao_preenchida.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:
        st.error(f"Ocorreu um erro durante a execução: {e}")
  
   

    
    

 
   

